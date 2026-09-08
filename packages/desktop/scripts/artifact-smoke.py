"""Offline structural verification for the bundled document toolkit (no model calls)."""
import json
from pathlib import Path
import sys
import prime_artifacts
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas
from PIL import Image
import pandas as pd
import matplotlib
matplotlib.use("Agg")
from matplotlib import pyplot as plt
out = Path(sys.argv[1]).resolve()
out.mkdir(parents=True, exist_ok=True)
assert prime_artifacts.status()["ready"]
doc = Document()
doc.add_heading("Prime Desktop 문서 검증", 0)
doc.add_paragraph("기본 도구로 생성하고 다시 읽은 문서입니다.")
doc.save(out / "report.docx")
book = Workbook()
book.active.append(["항목", "수량", "단가", "합계"])
book.active.append(["Sample", 3, 1200, "=B2*C2"])
book.save(out / "analysis.xlsx")
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[0])
slide.shapes.title.text = "Prime Desktop 검증"
slide.placeholders[1].text = "편집 가능한 발표자료"
deck.save(out / "presentation.pptx")
pdf = canvas.Canvas(str(out / "report.pdf"))
pdf.drawString(72, 740, "Prime Desktop artifact verification")
pdf.save()
pd.DataFrame({"category": ["A", "B"], "value": [3, 7]}).to_csv(out / "data.csv", index=False)
plt.bar(["A", "B"], [3, 7])
plt.ylabel("Count")
plt.tight_layout()
plt.savefig(out / "chart.png")
plt.close()
report = {"toolkit": prime_artifacts.status(), "artifacts": [prime_artifacts.inspect_file(out / name) for name in ("report.docx", "analysis.xlsx", "presentation.pptx", "report.pdf", "data.csv", "chart.png")], "modelRequests": 0}
assert "문서 검증" in report["artifacts"][0]["text"]
assert report["artifacts"][1]["sheets"][0]["rows"] == 2
assert len(report["artifacts"][2]["slides"]) == 1
assert report["artifacts"][3]["pages"] == 1
assert report["artifacts"][4]["data_rows"] == 2
(out / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps({"passed": True, "artifacts": len(report["artifacts"]), "report": str(out / "report.json"), "modelRequests": 0}))
