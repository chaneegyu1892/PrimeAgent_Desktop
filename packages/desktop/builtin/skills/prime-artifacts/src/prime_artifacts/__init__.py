"""Small, explicit diagnostics for Prime Desktop's document toolkit."""
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
import csv
import shutil

def status():
    """Report installed dependencies and optional renderers without installing anything."""
    packages = {}
    for name in ("python-docx", "openpyxl", "python-pptx", "pypdf", "reportlab", "pandas", "matplotlib", "Pillow"):
        try:
            packages[name] = version(name)
        except PackageNotFoundError:
            packages[name] = None
    return {"packages": packages, "ready": all(packages.values()), "renderers": {name: shutil.which(name) for name in ("soffice", "pdftoppm", "mutool", "tesseract")}}

def inspect_file(path):
    """Reopen an output artifact. Does not calculate formulas or render pages."""
    file = Path(path).expanduser().resolve(strict=True)
    result = {"path": str(file), "bytes": file.stat().st_size, "validation": "structure-only"}
    suffix = file.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        doc = Document(file)
        result.update(paragraphs=len(doc.paragraphs), tables=len(doc.tables), text="\n".join(p.text for p in doc.paragraphs)[:12000])
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(file, read_only=True, data_only=False)
        try:
            result["sheets"] = [{"name": s.title, "rows": s.max_row, "columns": s.max_column} for s in book]
        finally:
            book.close()
        result["formulas_recalculated"] = False
    elif suffix == ".pptx":
        from pptx import Presentation
        deck = Presentation(file)
        result["slides"] = [{"number": i+1, "text": "\n".join(s.text for s in slide.shapes if s.has_text_frame)[:3000], "overflow_shapes": sum(1 for s in slide.shapes if s.left < 0 or s.top < 0 or s.left+s.width > deck.slide_width or s.top+s.height > deck.slide_height)} for i, slide in enumerate(deck.slides)]
    elif suffix == ".pdf":
        from pypdf import PdfReader
        pdf = PdfReader(file)
        result.update(pages=len(pdf.pages), text="\n".join(page.extract_text() or "" for page in pdf.pages)[:12000])
    elif suffix == ".csv":
        with file.open(newline="", encoding="utf-8-sig") as handle:
            rows = csv.reader(handle)
            result["header"] = next(rows, [])
            result["data_rows"] = sum(1 for _ in rows)
    elif suffix in (".png", ".jpg", ".jpeg", ".webp"):
        from PIL import Image
        with Image.open(file) as image:
            result.update(width=image.width, height=image.height, format=image.format)
            image.verify()
    else:
        raise ValueError("Supported: DOCX, XLSX, PPTX, PDF, CSV, PNG, JPEG, WebP")
    return result
